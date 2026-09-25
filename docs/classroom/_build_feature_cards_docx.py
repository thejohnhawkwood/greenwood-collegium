"""Build a print-ready feature-cards.docx: two half-page cards, never split."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(__file__).with_name("feature-cards.docx")

INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x4A, 0x4A, 0x4A)
RULE = "1A1A1A"
CUT = "888888"
BODY = "Calibri"
MONO = "Consolas"

CARDS: list[dict[str, object]] = [
    {
        "num": "01",
        "title": "Four Ways to Find Things Out",
        "today": (
            "You learn things four different ways, and they are not the same job. look tells you what is in the room "
            "right now: people, objects, stuff on the ground, and exits. examine (or x) looks closer at one person or thing. "
            "talk is a private written reply from staff — only you see it — and can start a quest. say is something everyone "
            "in the room hears. Just looking around does not count as examining a clue. Staff stay put. Each one has one "
            "written conversation, not a choose-your-path chat."
        ),
        "play": [
            "In Lantern Court type look, examine noticeboard, examine porter, talk porter, then say a short line. Who sees or hears each one?",
            "Have your partner say while you talk. Then go to another room: look, examine its object, try talk on something that is not a person.",
            "talk flint or talk quill, then talk them again. What feels like a reminder, and what never changes?",
        ],
        "questions": [
            "Draw four boxes: quick glance vs a closer look, and only you vs everyone hears. Put look, examine, talk, and say in the boxes. Is a box empty? Missing feature, or a choice?",
            "In a game you know, how do you tell “just flavour” from “this is the clue”? Does Greenwood make that obvious, or do players guess?",
            "talk is private and written by the game. say is public and written by students. What goes wrong if those were one command? What goes wrong if Porter spoke with say like a student?",
            "Write one house rule for a new feature: what must show up in look, what should wait for examine, and what only talk is allowed to change (your quest — not the furniture everyone shares).",
            "When 30 people are looking and talking at once, which of these four ways of finding things out breaks first? What would you add that is not a private whisper?",
        ],
    },
    {
        "num": "02",
        "title": "Quests and What To Do Next",
        "today": (
            "Arrival starts by itself. The other three start only when you talk to the person who gives them "
            "(Quill, Tansy, or Alder). quests is a checklist: done and still to do. The clue objects stay in the world "
            "for everyone; only your progress is saved. If you examined a clue before you took the quest, you have to "
            "examine it again. Looking at a room is not the same as examining the clue. You only get the experience once."
        ),
        "play": [
            "Finish Arrival if you need to (look, say hello, take key, north), then type quests.",
            "talk quill (north, then west). Walk away. From memory, list what is left. Check with quests. Try to finish before you have all the clues.",
            "Examine one clue out of order. If you peeked before you accepted the quest, examine it again and watch quests.",
        ],
        "questions": [
            "Arrival teaches the first words. The others make you accept a job first. What does each style teach? If we added a fifth quest, which style would you use, and why?",
            "Copy one line from quests. Rewrite it two ways: a better hint, and a version that spoils the answer. What rule would keep a quest tracker between “too vague” and “too easy”?",
            "Having to examine again if you were early is a rule. Fair, or just extra work? How could the game remember “you already saw this” without letting people skip talking to the giver?",
            "Thirty students can examine the same blotter. What must never disappear from the room? What belongs only to you?",
            "Suggest a tiny pinned tracker (three lines max) that still works after you refresh. What should it show, and what should it hide so it does not spoil the mystery?",
        ],
    },
    {
        "num": "03",
        "title": "Knowing How You're Doing",
        "today": (
            "Level and experience show up as extra lines in the story (Arrival is about 10 experience and Level 2; "
            "later quests give 10 or 15, once). Fighting spends health and focus in that same scrolling story. "
            "There is no stats command and no always-on health bar. Your progress is still there after you refresh — "
            "but the numbers do not stay on screen."
        ),
        "play": [
            "After Arrival or finishing a quest, write your level and experience without scrolling. Then scroll and see if you were right.",
            "south, fight the dummy. After two actions, tell your partner your health, focus, and whether burning is going — no scrolling.",
            "Type quests. Does it help you decide whether to fight, rest, or keep exploring?",
        ],
        "questions": [
            "Name a real choice you made (fight, explore, stop, cast ember) that needed a number you did not have. What should have been on screen at that moment?",
            "Think of a game you know with health and mana bars that never leave the screen. What from that game would help here? What would make Greenwood feel like a generic clicker RPG?",
            "Would you rather type stats when you need it, or keep a small status strip always visible? Your idea must still work after refresh and in a packed courtyard.",
            "If you add bars, what words sit next to them so colour and motion are never the only way to read them?",
            "From playing, list numbers the game already knows (health, focus, level, experience…). Your feature should show those — not invent new ones the game does not have.",
        ],
    },
    {
        "num": "04",
        "title": "Combat You Can Follow",
        "today": (
            "South Orchard has a practice dummy. Fights take turns. You type; the game decides hits and damage — your screen "
            "only shows the result. attack dummy starts it; later you can type attack. cast ember dummy spends focus and can "
            "leave burning. If you lose, you wake in the Infirmary and keep your items. Flint will explain the words; talking "
            "does not change the fight. There is no health frame and no list of legal moves on screen."
        ),
        "play": [
            "south, talk flint, examine dummy, attack dummy.",
            "Mix attack and cast ember. After each line, jot: whose turn, what you are allowed to type, health or focus if it said, any burning.",
            "Have your partner say in the orchard while you fight. Optional: lose on purpose and read the Infirmary arrival as “what the game did,” not just story.",
        ],
        "questions": [
            "Sketch the fight as a sequence of steps: start, your turn, result, dummy’s turn, win or Infirmary. Where did the scrolling story hide the step you were on?",
            "Ember costs focus and can add burning. What would make it feel different from a normal attack — without a flashy effect deciding the damage?",
            "A lot of games take over the whole screen for combat. Here it is more story lines. Suggest a feature: what belongs in a small fight box, and what must stay in the story so someone who looks away can still reconstruct the fight?",
            "Waking in the Infirmary with your stuff is a class-fair rule. What does a player no longer have to worry about? Is that the right trade for a period?",
            "Chat and fighting share one story. Suggest a rule for mixing them so the room can still talk and you do not miss a number you needed.",
        ],
    },
    {
        "num": "05",
        "title": "Your Pack and Who Owns What",
        "today": (
            "take, drop, and inventory (or i) are how you handle stuff. Each Collegian gets their own Small Copper Key. "
            "The moss-bound primer in the Library Stacks is one book for the whole school — two people cannot both have it. "
            "The noticeboard and staff are scenery: you cannot pick them up. You cannot wear gear yet, and there is no bag picture."
        ),
        "play": [
            "take key, inventory, drop key, look, take key. Then try take on the noticeboard and on Porter. What does the refusal teach you?",
            "Race your partner to the Library (north, west) and both try take primer. One of you should fail. Why?",
            "Drop the key in a room, leave, come back. Who is allowed to pick it up?",
        ],
        "questions": [
            "The game already has two kinds of stuff: “everyone gets their own” (the key) and “only one in the school” (the primer). When should a new item be which? How should the game tell you before you fail?",
            "You cannot wear gear. Would an equip feature change a choice you make today, or is it just a costume? What would fighting or quests have to gain before worn gear matters?",
            "Write two rules for a new feature: one a player can read (“you can pick up…”) and one the game follows (“take fails when…”). Use the noticeboard vs the key as your examples.",
            "If you drop the key and a classmate takes it, is that part of the game, griefing, or a missing lock? Suggest a class-safe rule that is not “nobody may drop items.”",
            "Suggest a bag view that still needs typed take, drop, and inventory. How would you group things (quest, carry, junk) so it helps after a long period?",
        ],
    },
    {
        "num": "06",
        "title": "Finding Your Way",
        "today": (
            "The school is 25 rooms joined by north / south / east / west — rooms and the paths between them, not an open "
            "world you walk pixel by pixel. look lists the exits from the room you are in. There is no map command and no "
            "picture of where you have been. The game file has map positions; players never see them. Hidden doors are not in play yet."
        ),
        "play": [
            "From Lantern Court, reach the Clock Tower with no notes (north, north, east). Come back a different way.",
            "From memory, sketch every room you walked and the exits you are sure of. Mark guesses with ?.",
            "Find the Quiet Chapel or the Archive Cellar without this card. When you are lost, you may only use look. Compare sketches with your partner.",
        ],
        "questions": [
            "A later map would be rooms and the paths between them — not a pretty painting. What would a small map show that look does not? When would pretty art lie?",
            "Like fog of war in other games: do you show only rooms you have visited, or also exits you have already seen? What about a room you walked through without looking?",
            "Name a map habit from a game you know (mini-map, compass, fast travel). Which would help here? Which would make typing north / south pointless?",
            "Thirty students will get lost in the same 25 rooms. Is that a writing problem, a missing map, or part of the first-hour adventure? Defend one.",
            "Suggest an honest map feature: you, rooms you have found, exits you have found, no secrets, readable if someone prints it in black and white. What must the game remember after you log out?",
        ],
    },
    {
        "num": "07",
        "title": "Reading the Live Story",
        "today": (
            "The gold window is the whole game: speech, people arriving, looks, fights, and quests all share one scrolling "
            "story. Colours and written labels mark different kinds of line. If you scroll up to reread, new lines stop "
            "chasing you; New messages jumps you back. You cannot filter or pin a line."
        ),
        "play": [
            "Stay in a busy Lantern Court. look, then stand still for two minutes. How many kinds of line show up?",
            "Scroll up to reread a look and stay there while people keep arriving and talking. Use New messages.",
            "Start a quest update or a fight, then immediately say. Which line grabs you? Tell your partner what happened in the last minute if they looked away.",
        ],
        "questions": [
            "This is not a chat app with a game stuck on. The story is the game. Which new lines should yank you to the bottom, and which should wait behind New messages?",
            "Labels exist so colour is not the only clue. Suggest better labels that still work on the projector and if someone prints the page in black and white.",
            "Think of a fight log in a game you know. What “show only…” or “keep this line” feature would help here — without hiding a finished quest or a teacher mute?",
            "If the screen split (story + one extra panel), what is the one thing that must never leave the story, and why?",
            "Suggest a feature, then say how you would know it worked: after ten crowded minutes, what would you watch or count?",
        ],
    },
    {
        "num": "08",
        "title": "The Words You Type",
        "today": (
            "You play by typing words. help lists every word the game understands today; help look explains one word. "
            "Up and down arrows bring back lines you already typed. A typo gets the same “Did you mean” list every time — "
            "the game does not guess your closest word. Shortcuts exist (n/s/e/w, x, i). There is no tab-complete and no "
            "hint bar. Words from other games (map, stats, who, equip) are not here yet."
        ),
        "play": [
            "Type help, then help talk, help cast, help quests. Teach your partner one word using only what help printed.",
            "Misspell on purpose (attak dummy, lok, invntry). Then type words other games have that this one does not: use, map, stats, who, equip.",
            "During a fight or a quest, notice which words you type from memory and which you look up. Use the up arrow once to go faster.",
        ],
        "questions": [
            "help is a word list, not a coach. Suggest a feature that hints what to type next without turning the game into buttons.",
            "The “Did you mean” list never changes. When is that more honest than auto-correct? When does it waste a turn in a fight?",
            "Players bring words from other games. Which missing words are a real hole today (the game already does that job)? Which would be a lie because that job is not built yet?",
            "Suggest shortcuts and one “finish the word for me” idea that still works if someone never uses a mouse and never relies on colour.",
            "Write one example for your idea: what the player types, what the game should say back, and one guess the game must not make.",
        ],
    },
]

COMMANDS = [
    "help look",
    "help talk",
    "help cast",
    "help quests",
    "attak dummy",
    "invntry",
    "New messages",
    "examine noticeboard",
    "examine porter",
    "examine dummy",
    "talk porter",
    "talk flint",
    "talk quill",
    "talk tansy",
    "take key",
    "drop key",
    "take primer",
    "say hello",
    "attack dummy",
    "cast ember dummy",
    "cast ember",
    "north, north, east",
    "inventory",
    "quests",
    "examine",
    "look",
    "help",
    "talk",
    "say",
    "take",
    "drop",
    "north",
    "south",
    "east",
    "west",
    "stats",
    "map",
    "who",
    "equip",
    "use",
    "lok",
    "n s e w",
    "x porter",
]


def _set_run_font(run, name: str, size_pt: float, *, bold: bool = False, color: RGBColor = INK) -> None:
    run.bold = bold
    run.font.size = Pt(size_pt)
    run.font.color.rgb = color
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:cs"), name)


def _shade(cell, fill: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def _set_cell_margins(cell, *, top: int, bottom: int, left: int, right: int) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for edge, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def _set_cell_borders(cell, **edges: tuple[str, str, str]) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        if edge in edges:
            val, sz, color = edges[edge]
            node.set(qn("w:val"), val)
            node.set(qn("w:sz"), sz)
            node.set(qn("w:space"), "0")
            node.set(qn("w:color"), color)
        else:
            node.set(qn("w:val"), "nil")
        tcBorders.append(node)
    tcPr.append(tcBorders)


def _cant_split(row) -> None:
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))


def _row_height_exact(row, inches: float) -> None:
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    trHeight = OxmlElement("w:trHeight")
    trHeight.set(qn("w:val"), str(int(inches * 1440)))
    trHeight.set(qn("w:hRule"), "exact")
    trPr.append(trHeight)


def _keep_together(paragraph) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    for tag in ("keepNext", "keepLines", "widowControl"):
        el = OxmlElement(f"w:{tag}")
        if tag == "widowControl":
            el.set(qn("w:val"), "true")
        pPr.append(el)


def _add_mixed(paragraph, text: str, size: float, *, bold: bool, color: RGBColor, mono_words: bool) -> None:
    if not mono_words:
        run = paragraph.add_run(text)
        _set_run_font(run, BODY, size, bold=bold, color=color)
        return
    pattern = "(" + "|".join(re.escape(cmd) for cmd in sorted(COMMANDS, key=len, reverse=True)) + ")"
    parts = re.split(pattern, text)
    for part in parts:
        if not part:
            continue
        if part in COMMANDS:
            run = paragraph.add_run(part)
            _set_run_font(run, MONO, size - 0.5, bold=True, color=color)
        else:
            run = paragraph.add_run(part)
            _set_run_font(run, BODY, size, bold=bold, color=color)


def _add_p(cell, *, before: float, after: float, leading: float, indent: float = 0):
    p = cell.paragraphs[0] if not cell.paragraphs[0].text and not cell.paragraphs[0].runs else cell.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p.paragraph_format.line_spacing = Pt(leading)
    p.paragraph_format.left_indent = Inches(indent)
    _keep_together(p)
    return p


def _fill_card(cell, card: dict[str, object], *, cut_hint: bool, total: int) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    _set_cell_margins(cell, top=60, bottom=40, left=90, right=90)
    _shade(cell, "F7F4EC")
    _set_cell_borders(
        cell,
        top=("single", "12", RULE),
        left=("single", "12", RULE),
        bottom=("dashSmallGap" if cut_hint else "single", "8" if cut_hint else "12", CUT if cut_hint else RULE),
        right=("single", "12", RULE),
    )

    eyebrow = _add_p(cell, before=0, after=0, leading=10)
    _set_run_font(eyebrow.add_run("THE GREENWOOD COLLEGIUM"), BODY, 8, bold=True, color=MUTED)
    _set_run_font(eyebrow.add_run("   ·   FEATURE CARD   "), BODY, 8, color=MUTED)
    _set_run_font(eyebrow.add_run(f"{card['num']} / {total:02d}"), BODY, 8, bold=True)

    title = _add_p(cell, before=1, after=1, leading=16)
    _set_run_font(title.add_run(str(card["title"])), BODY, 14, bold=True)

    rule = _add_p(cell, before=0, after=2, leading=3)
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "10")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), RULE)
    pBdr.append(bottom)
    rule._p.get_or_add_pPr().append(pBdr)

    how = _add_p(cell, before=0, after=3, leading=10)
    _set_run_font(
        how.add_run("Play hard for 10 minutes. Then suggest a feature: who it helps, what you give up, how you’d know it worked."),
        BODY,
        8,
        color=MUTED,
    )

    label = _add_p(cell, before=0, after=0, leading=11)
    _set_run_font(label.add_run("What it is today"), BODY, 9, bold=True)

    today = _add_p(cell, before=0, after=3, leading=10.5)
    _add_mixed(today, str(card["today"]), 8.5, bold=False, color=INK, mono_words=True)

    play_l = _add_p(cell, before=1, after=0, leading=11)
    _set_run_font(play_l.add_run("Try this first"), BODY, 9, bold=True)

    for i, step in enumerate(card["play"], start=1):  # type: ignore[arg-type]
        p = _add_p(cell, before=0, after=0, leading=10.5, indent=0.14)
        _add_mixed(p, f"{i}. {step}", 8.5, bold=False, color=INK, mono_words=True)

    q_l = _add_p(cell, before=3, after=0, leading=11)
    _set_run_font(q_l.add_run("Suggest and improve"), BODY, 9, bold=True)

    for i, q in enumerate(card["questions"], start=1):  # type: ignore[arg-type]
        p = _add_p(cell, before=0, after=1, leading=10.5, indent=0.14)
        _add_mixed(p, f"{i}. {q}", 8.5, bold=False, color=INK, mono_words=True)

    foot = _add_p(cell, before=3, after=0, leading=9)
    _set_run_font(
        foot.add_run(
            "Keep typed commands  ·  don’t use colour as the only clue  ·  the game decides what happened  ·  no whispers, PvP, or real names"
        ),
        BODY,
        7.5,
        color=MUTED,
    )

    if cut_hint:
        cut = _add_p(cell, before=2, after=0, leading=9)
        cut.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_run_font(cut.add_run("- - -  cut along this line  - - -"), BODY, 7.5, color=MUTED)


def _suppress_table_borders(table) -> None:
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        borders.append(node)
    tblPr.append(borders)


def _apply_page(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)
    section.top_margin = Inches(0.35)
    section.bottom_margin = Inches(0.35)
    section.header_distance = Inches(0.2)
    section.footer_distance = Inches(0.2)


def _zero_trailing_paragraph(paragraph) -> None:
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = Pt(1)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY


def _set_table_width(table, width_dxa: int) -> None:
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(width_dxa))
    tblW.set(qn("w:type"), "dxa")
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)


def _clear_body_paragraphs(doc) -> None:
    body = doc.element.body
    for child in list(body):
        if child.tag == qn("w:p"):
            body.remove(child)


def build() -> Path:
    doc = Document()
    _apply_page(doc.sections[0])
    _clear_body_paragraphs(doc)

    normal = doc.styles["Normal"]
    normal.font.name = BODY
    normal.font.size = Pt(9)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = Pt(11)

    usable_width = Inches(7.5)
    width_dxa = 10800
    card_height = 4.7
    total = len(CARDS)

    for page_index in range(0, total, 2):
        table = doc.add_table(rows=2, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        _suppress_table_borders(table)
        _set_table_width(table, width_dxa)
        table.columns[0].width = usable_width

        for row_i, card in enumerate((CARDS[page_index], CARDS[page_index + 1])):
            row = table.rows[row_i]
            _cant_split(row)
            _row_height_exact(row, card_height)
            cell = row.cells[0]
            cell.width = usable_width
            _fill_card(cell, card, cut_hint=(row_i == 0), total=total)

        if page_index + 2 < total:
            br = doc.add_paragraph()
            _zero_trailing_paragraph(br)
            br.add_run().add_break(WD_BREAK.PAGE)

    trailing = doc.add_paragraph()
    _zero_trailing_paragraph(trailing)

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(path)
